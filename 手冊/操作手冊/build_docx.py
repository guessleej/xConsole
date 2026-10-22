#!/usr/bin/env python3
"""xConsole 功能說明：md 為真本 → docx。

用法：python3 build_docx.py
流程：先產一次 docx → soffice 轉 PDF → 量出每個章節實際在第幾頁 → 回填目錄頁碼 → 再產一次。
目錄用書籤＋內部連結（可點），不用 Word 的 TOC 功能變數（別的預覽工具會顯示空白）。
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "xConsole主要功能說明.md"
OUT = HERE / "xConsole主要功能說明.docx"

RED = RGBColor(0xBF, 0x18, 0x1F)
INK = RGBColor(0x16, 0x20, 0x2A)
INK2 = RGBColor(0x3D, 0x4B, 0x56)
GREY = RGBColor(0x6B, 0x7A, 0x85)
# Segoe UI 在 Mac 上沒有，Word／LibreOffice 都會退回 Times 這類襯線字；英數字一律跟中文同一套
FONT_ZH, FONT_EN, FONT_CODE = "微軟正黑體", "Microsoft JhengHei", "Consolas"


# ------------------------------------------------------------------ md 解析

def parse(md):
    meta, body = {}, md
    m = re.match(r"^---\n(.*?)\n---\n", md, re.S)
    if m:
        for line in m.group(1).splitlines():
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
        body = md[m.end():]
    blocks, lines, i = [], body.splitlines(), 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            i += 1
            continue
        if ln.startswith("# "):
            blocks.append(("h1", ln[2:].strip()))
        elif ln.startswith("## "):
            blocks.append(("h2", ln[3:].strip()))
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            cap = None
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and lines[i].startswith(": "):
                cap = lines[i][2:].strip()
                i += 1
            blocks.append(("table", rows, cap))
            continue
        else:
            para = [ln]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", ": ")):
                para.append(lines[i].strip())
                i += 1
            blocks.append(("p", "".join(para)))
            continue
        i += 1
    return meta, blocks


# ------------------------------------------------------------------ 樣式工具

def set_font(run, size=None, bold=None, color=None, code=False):
    f = FONT_CODE if code else FONT_EN
    run.font.name = f
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), f)
    rfonts.set(qn("w:eastAsia"), FONT_ZH)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=10.5, color=INK, bold=False):
    """處理 **粗體** 與 `程式碼`。"""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            set_font(par.add_run(tok[2:-2]), size, True, color)
        elif tok.startswith("`"):
            set_font(par.add_run(tok[1:-1]), size - 0.5, bold, INK2, code=True)
        else:
            set_font(par.add_run(tok), size, bold, color)


def spacing(par, before=0, after=6, line=1.5):
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


_bm_id = [0]


def bookmark(par, name):
    _bm_id[0] += 1
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), str(_bm_id[0]))
    s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), str(_bm_id[0]))
    par._p.insert(0, s)
    par._p.append(e)


def link(par, anchor, text, size, bold=False, color=INK):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    h.set(qn("w:history"), "1")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), FONT_EN)
    fonts.set(qn("w:eastAsia"), FONT_ZH)
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), str(color))
    rpr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def vcenter(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tcpr.append(v)


def cell_margins(table, top=70, bottom=70, left=110, right=110):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def borders(table):
    """只留列與列之間的細分隔線（功能性），不要直線、不要外框。"""
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if side in ("insideH", "bottom"):
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "DDE3E8")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)


def field(par, code, size=8.5, color=GREY):
    r1 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {code} "
    r2._r.append(it)
    r3 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r3._r.append(fc)
    for r in (r1, r2, r3):
        set_font(r, size, False, color)


# ------------------------------------------------------------------ 建置

def anchor_of(text):
    import hashlib
    return "h_" + hashlib.md5(text.encode()).hexdigest()[:10]


def build(meta, blocks, pages):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.left_margin = sec.right_margin = Cm(2.3)

    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_ZH)

    # ---------- 封面：眉標 → 主標題 → 副標語 → 資訊表 → 公司與日期
    # 封面高度用段落間距控制，不要用空段落堆——空行累積起來會把日期擠到第二頁（實際踩到）
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta["eyebrow"]), 12, True, RED); spacing(p, 150, 10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta["title"]), 30, True, INK); spacing(p, 0, 12, 1.2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta["subtitle"]), 13, False, INK2); spacing(p, 0, 60)

    info = [("文件名稱", meta["title"] + "　主要功能說明"), ("版本", meta["version"]),
            ("發行日期", meta["date"]), ("編製單位", meta["unit"]),
            ("聯絡窗口", meta["contact"]), ("文件密級", meta["classification"])]
    t = doc.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t); cell_margins(t)
    for i, (k, v) in enumerate(info):
        a, b = t.rows[i].cells
        a.width, b.width = Cm(3.4), Cm(10.6)
        shade(a, "F4F7F9")
        pa = a.paragraphs[0]; set_font(pa.add_run(k), 10, True, INK2)
        pb = b.paragraphs[0]; set_font(pb.add_run(v), 10, False, INK)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("云碩科技股份有限公司"), 11, True, INK); spacing(p, 90, 2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta["date"]), 9.5, False, GREY); spacing(p, 0, 0)
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---------- 目錄（靜態表格、每列可點到對應章節）
    p = doc.add_paragraph()
    set_font(p.add_run("目錄"), 18, True, INK); spacing(p, 0, 14)
    heads = [b for b in blocks if b[0] in ("h1", "h2")]
    toc = doc.add_table(rows=len(heads), cols=2)
    borders(toc); cell_margins(toc, 22, 22, 60, 60)
    for i, (kind, text) in enumerate(heads):
        a, b = toc.rows[i].cells
        a.width, b.width = Cm(14.2), Cm(2.2)
        vcenter(a); vcenter(b)
        pa = a.paragraphs[0]; spacing(pa, 0, 0, 1.15)
        if kind == "h2":
            pa.paragraph_format.left_indent = Cm(0.8)
        link(pa, anchor_of(text), text, 10.5 if kind == "h1" else 10,
             bold=(kind == "h1"), color=INK if kind == "h1" else INK2)
        pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT; spacing(pb, 0, 0, 1.15)
        set_font(pb.add_run(str(pages.get(text, "—"))), 10, kind == "h1", INK2)

    # ---------- 正文
    for blk in blocks:
        kind = blk[0]
        if kind == "h1":
            # 每一章從新的一頁開始：用標題自己的「段前分頁」，不要另放一段分頁符——
            # 上一章剛好寫到頁底時，那個分頁段落會自己被擠到新頁再分一次，多出一張白紙（實際踩到）
            h = doc.add_paragraph()
            h.paragraph_format.page_break_before = True
            set_font(h.add_run(blk[1]), 18, True, INK); spacing(h, 0, 14, 1.25)
            bookmark(h, anchor_of(blk[1]))
            h.paragraph_format.keep_with_next = True
        elif kind == "h2":
            h = doc.add_paragraph()
            set_font(h.add_run(blk[1]), 13, True, RED); spacing(h, 14, 6, 1.25)
            bookmark(h, anchor_of(blk[1]))
            h.paragraph_format.keep_with_next = True
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(21)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, blk[1]); spacing(p, 0, 8, 1.6)
        elif kind == "table":
            rows, cap = blk[1], blk[2]
            if cap:
                c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(c.add_run(cap), 9.5, True, INK2); spacing(c, 8, 4, 1.2)
                c.paragraph_format.keep_with_next = True
            ncol = len(rows[0])
            tb = doc.add_table(rows=len(rows), cols=ncol)
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(tb); cell_margins(tb)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = tb.rows[ri].cells[ci]
                    txt = row[ci] if ci < len(row) else ""
                    par = cell.paragraphs[0]; spacing(par, 0, 0, 1.3); vcenter(cell)
                    if ri == 0:
                        shade(cell, "BF181F")
                        set_font(par.add_run(txt), 9.5, True, RGBColor(0xFF, 0xFF, 0xFF))
                    else:
                        add_inline(par, txt, 9.5, INK)
                if ri == 0:
                    repeat_header(tb.rows[0])
                # 每一列本身都不准拆成兩半掉到下一頁
                trpr = tb.rows[ri]._tr.get_or_add_trPr()
                if trpr.find(qn("w:cantSplit")) is None:
                    trpr.append(OxmlElement("w:cantSplit"))
                # 短表格整張不拆頁：三列的表最後一列單獨掉到下一頁很難看（表 1-1 實際踩到）。
                # 門檻不要開太大：表格整張推到下一頁時，章標題會孤零零留在上一頁（試過 12 列，第五章就這樣）。
                # 表格太長就該把內容收斂，不是硬撐門檻
                if len(rows) <= 10 and ri < len(rows) - 1:
                    for cell in tb.rows[ri].cells:
                        for par in cell.paragraphs:
                            par.paragraph_format.keep_with_next = True
            sp = doc.add_paragraph(); spacing(sp, 0, 6, 1.0)

    # ---------- 結尾署名
    p = doc.add_paragraph(); spacing(p, 30, 0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"本文件由{meta['unit']}編製　{meta['date'][:4]} 年 {int(meta['date'][5:7])} 月"),
             9.5, False, GREY)

    # ---------- 頁尾（封面不顯示）
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run("云碩科技股份有限公司　機密文件　第 "), 8.5, False, GREY)
    field(fp, "PAGE")
    set_font(fp.add_run(" 頁，共 "), 8.5, False, GREY)
    field(fp, "NUMPAGES")
    set_font(fp.add_run(" 頁"), 8.5, False, GREY)

    doc.save(OUT)


# ------------------------------------------------------------------ 頁碼回填

def to_pdf(docx):
    out = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", str(docx), "--outdir", str(out)],
                   check=True, capture_output=True)
    return out / (docx.stem + ".pdf")


def measure_pages(pdf, heads):
    txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
    pages = txt.split("\f")
    # 目錄頁本身就印著所有章節名，要跳過它（目錄可能不只一頁）
    last_toc = max((i for i, pg in enumerate(pages[:4]) if "目錄" in pg), default=1)
    start = last_toc + 1
    while start < len(pages) and sum(re.sub(r"\s+", "", t) in re.sub(r"\s+", "", pages[start])
                                     for _, t in heads) > len(heads) // 2:
        start += 1
    found = {}
    for kind, text in heads:
        key = re.sub(r"\s+", "", text)
        for pi in range(start, len(pages)):
            if key in re.sub(r"\s+", "", pages[pi]):
                found[text] = pi + 1
                start = pi
                break
    return found, len(pages) - (1 if pages and not pages[-1].strip() else 0)


def main():
    meta, blocks = parse(SRC.read_text(encoding="utf-8"))
    heads = [b for b in blocks if b[0] in ("h1", "h2")]
    pages = {}
    for rnd in (1, 2, 3):
        build(meta, blocks, pages)
        pdf = to_pdf(OUT)
        new, total = measure_pages(pdf, heads)
        if new == pages:
            break
        pages = new
    missing = [t for _, t in heads if t not in pages]
    print(f"已產出 {OUT.name}，共 {total} 頁；目錄 {len(heads) - len(missing)}/{len(heads)} 項回填頁碼")
    if missing:
        print("找不到頁碼：", missing)
        return 1
    pdf_out = HERE / (OUT.stem + ".pdf")
    pdf_out.write_bytes(pdf.read_bytes())
    print(f"PDF：{pdf_out.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
