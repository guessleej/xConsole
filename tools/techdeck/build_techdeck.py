# -*- coding: utf-8 -*-
"""xConsole 模型主控台：給客戶技術長看的技術簡報（pptx＋pdf＋總覽圖）。

    python build_techdeck.py                 # 重畫新圖、產 pptx、轉 pdf、出總覽圖
    python build_techdeck.py --skip-diagrams --no-pdf

骨架照 xcloud-deck-build 第十節（產品介紹說明、技術演進敘事版）：每頁一句 goldline 金句收尾，
備忘稿寫「講法導引」與「口徑註記」。版型 helper 與 QA 檢查沿用 tools/saleskit/build_saleskit.py 的寫法
（複製一份到這裡，不 import，避免互相牽動）。每個數字都對得到 手冊/資安證據、docs 與程式碼。
執行要用有 python-pptx 的 Python（本機是 /opt/homebrew/bin/python3.14）。
"""
import argparse
import math
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches as I, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, os.path.expanduser("~/.claude/skills/arch-diagram-style"))
sys.path.insert(0, os.path.expanduser("~/.claude/skills/xcloud-deck-build/assets"))
from msicon_png import icon_png  # noqa: E402
from pptx_anim import add_layers, animate_slide, check_timing  # noqa: E402

ROOT = HERE.parents[1]                                   # code/
KIT = ROOT / "手冊" / "Sales Kit"
OUTDIR = KIT / "技術長版"
FIG = OUTDIR / "圖"
OLDFIG = KIT / "圖"                                      # Sales Kit 共用圖（全景、GPU）與截圖
SHOT = OLDFIG / "截圖"
QA = OUTDIR / "QA"
TPL = Path(os.path.expanduser("~/.claude/skills/xcloud-deck-build/assets/xcloud-template.pptx"))
SOFFICE = "/opt/homebrew/bin/soffice"
FONT = "Microsoft JhengHei"
DATE = "2026年10月"
NAME = "xConsole-技術簡報"

RED = RGBColor(0xBF, 0x18, 0x1F); RED_D = RGBColor(0x9C, 0x12, 0x18); RTINT = RGBColor(0xFB, 0xEB, 0xEC)
CHAR = RGBColor(0x2D, 0x2D, 0x2D); SLATE = RGBColor(0x4A, 0x4A, 0x4A); GREY = RGBColor(0x6E, 0x6E, 0x6E)
LGREY = RGBColor(0xF1, 0xF2, 0xF4); WHITE = RGBColor(0xFF, 0xFF, 0xFF); STEEL = RGBColor(0x5B, 0x6B, 0x79)
GREEN = RGBColor(0x3E, 0x8E, 0x54); LINEC = RGBColor(0xD9, 0xDC, 0xE0)

STATUS = {"done": ("已實作且有測試", GREEN, WHITE), "verify": ("已實作待實機驗證", STEEL, WHITE),
          "plan": ("規劃中", LINEC, CHAR)}

_CJK = r"[㐀-鿿　-〿＀-￯]"
PROBLEMS = []
KEEP_GOING = False


def nz(s):
    """中文與英數之間不留空白（使用者規定）。"""
    s = re.sub(rf"({_CJK}) +([A-Za-z0-9])", r"\1\2", str(s))
    return re.sub(rf"([A-Za-z0-9%]) +({_CJK})", r"\1\2", s)


def text_in(s, size):
    """估算字串寬度（吋）：全形一字寬，半形約 0.56 字寬。"""
    w = 0.0
    for ch in s:
        o = ord(ch)
        if ch == " ":
            w += 0.3
        elif 0x2E80 <= o <= 0x9FFF or 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFF60:
            w += 1.0
        else:
            w += 0.56
    # 中文與英數交界處，PowerPoint 與 LibreOffice 都會自動補約四分之一字寬的間距（東亞自動間距），要算進去
    w += 0.25 * len(re.findall(rf"(?<={_CJK})[A-Za-z0-9]|(?<=[A-Za-z0-9]){_CJK}", s))
    return w * size / 72 * 1.08      # 實際斷行比估算寬，留餘裕


def need_height(paras, size, width, sp):
    lines = sum(max(1, math.ceil(text_in(p, size) / max(width, 0.1) - 0.02)) for p in paras)
    return lines * size * sp * 1.2 / 72


class Deck:
    def __init__(self):
        self.prs = Presentation(str(TPL))
        for sid in list(self.prs.slides._sldIdLst):
            try:
                self.prs.part.drop_rel(sid.get(qn("r:id")))
            except Exception:
                pass
            self.prs.slides._sldIdLst.remove(sid)
        self.lay = {l.name: l for l in self.prs.slide_layouts}
        self.icons = {}

    # 基本元件
    def add(self, name="內容"):
        s = self.prs.slides.add_slide(self.lay[name])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        return s

    @staticmethod
    def _font(run, size, bold, color):
        run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color; run.font.name = FONT
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {}); rPr.append(e)
            e.set("typeface", FONT)

    def tx(self, s, l, t, w, h, runs, size=14, color=SLATE, bold=False, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP, sp=1.0, check=True, space_after=0):
        tb = s.shapes.add_textbox(I(l), I(t), I(w), I(h)); tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        if isinstance(runs, str):
            runs = [(runs, color, bold)]
        p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = sp
        if space_after:
            p.space_after = Pt(space_after)
        paras, cur = [], ""
        for item in runs:
            if item == "\n":
                p = tf.add_paragraph(); p.alignment = align; p.line_spacing = sp
                if space_after:
                    p.space_after = Pt(space_after)
                paras.append(cur); cur = ""; continue
            t_, c_, b_ = item if isinstance(item, tuple) else (item, color, bold)
            for k, part in enumerate(nz(t_).split("\n")):
                if k:
                    p = tf.add_paragraph(); p.alignment = align; p.line_spacing = sp
                    if space_after:
                        p.space_after = Pt(space_after)
                    paras.append(cur); cur = ""
                if part:
                    r = p.add_run(); r.text = part
                    self._font(r, size, b_ if b_ is not None else bold, c_ if c_ is not None else color)
                    cur += part
        paras.append(cur)
        if check:
            need = need_height(paras, size, w, sp) + len(paras) * space_after / 72
            if need > h + 0.05:
                PROBLEMS.append(f"第{len(self.prs.slides)}頁 文字可能溢出（需約{need:.2f}吋，框高{h:.2f}吋）：{''.join(paras)[:24]}")
        return tb

    def rr(self, s, l, t, w, h, fill, rad=0.06, name=None):
        sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(l), I(t), I(w), I(h))
        try:
            sp.adjustments[0] = rad
        except Exception:
            pass
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
        sp.line.fill.background(); sp.shadow.inherit = False
        if name:
            sp.name = name
        return sp

    def rect(self, s, l, t, w, h, fill, name="cover-mask"):
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(l), I(t), I(w), I(h))
        sp.fill.solid(); sp.fill.fore_color.rgb = fill; sp.line.fill.background(); sp.shadow.inherit = False
        sp.name = name
        return sp

    def oval(self, s, l, t, d, fill):
        sp = s.shapes.add_shape(MSO_SHAPE.OVAL, I(l), I(t), I(d), I(d))
        sp.fill.solid(); sp.fill.fore_color.rgb = fill; sp.line.fill.background(); sp.shadow.inherit = False
        return sp

    def icon(self, s, name, x, y, color="#BF181F", size=0.30):
        key = (name, color)
        if key not in self.icons:
            self.icons[key] = icon_png(name, color, 128)
        s.shapes.add_picture(self.icons[key], I(x), I(y), width=I(size))

    def title(self, s, t, sub=None):
        self.tx(s, 2.5, 0.42, 10.2, 0.6, t, size=24, color=CHAR, bold=True, anchor=MSO_ANCHOR.MIDDLE).name = "xh-title"
        if sub:
            self.tx(s, 2.5, 1.02, 10.2, 0.32, sub, size=13, color=GREY).name = "xh-title"

    @staticmethod
    def notes(s, guide, caliber):
        tf = s.notes_slide.notes_text_frame
        tf.text = "講法導引：" + nz(guide)
        p = tf.add_paragraph(); p.text = "口徑註記：" + nz(caliber)

    # 版型
    def goldline(self, s, text, y=6.14):
        """紅暈色底金句條，每頁一句；寬 10.4 吋避開右下 logo。"""
        h = max(0.54, need_height([nz(text)], 15, 9.9, 1.05) + 0.16)
        self.rr(s, 0.68, y, 10.4, h, RTINT, rad=0.14, name="goldline")
        self.tx(s, 0.92, y, 9.95, h, text, size=15, color=RED_D, bold=True, anchor=MSO_ANCHOR.MIDDLE, sp=1.05)
        if y + h > 6.80:
            PROBLEMS.append(f"第{len(self.prs.slides)}頁 金句條太低：{text[:16]}")

    def status(self, s, y, items, x=0.68, w=11.9):
        """模組狀態列：狀態膠囊＋一句說明，一列一項。回傳下一列的 y。"""
        for kind, note in items:
            label, fill, fg = STATUS[kind]
            pw = text_in(label, 12.5) + 0.34
            self.rr(s, x, y, pw, 0.32, fill, rad=0.5, name="status")
            self.tx(s, x, y, pw, 0.32, label, size=12.5, color=fg, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            nh = max(0.32, need_height([nz(note)], 12.5, w - pw - 0.16, 1.05))
            if nh > 0.33:
                PROBLEMS.append(f"第{len(self.prs.slides)}頁 狀態說明超過一行：{note[:20]}")
            self.tx(s, x + pw + 0.16, y + (0.02 if nh <= 0.33 else 0.02), w - pw - 0.16, nh, note, size=12.5, color=GREY,
                    anchor=MSO_ANCHOR.MIDDLE if nh <= 0.33 else MSO_ANCHOR.TOP, sp=1.05)
            y += max(0.32, nh) + 0.10
        return y

    def status_inline(self, s, items, x=2.5, y=1.02, w=10.2):
        """圖頁用：狀態膠囊與說明排在標題下方同一列（取代副標），讓圖可以放到約12吋寬。"""
        cx = x
        for kind, note in items:
            label, fill, fg = STATUS[kind]
            pw = text_in(label, 12.5) + 0.30
            self.rr(s, cx, y, pw, 0.30, fill, rad=0.5, name="status")
            self.tx(s, cx, y, pw, 0.30, label, size=12.5, color=fg, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            nw = text_in(nz(note), 12.5) + 0.05
            self.tx(s, cx + pw + 0.10, y, nw, 0.30, note, size=12.5, color=GREY, anchor=MSO_ANCHOR.MIDDLE, check=False)
            cx += pw + 0.10 + nw + 0.25
        if cx - 0.25 > x + w + 0.3:
            PROBLEMS.append(f"第{len(self.prs.slides)}頁 狀態列太長（到{cx:.2f}吋）")

    def paras(self, s, x, y, w, items, size=13, sp=1.16, gap=0.10, lead_color=CHAR):
        """機制說明段落：每段「粗體引題＋敘述」。回傳結束 y。"""
        for item in items:
            lead, body = item if isinstance(item, tuple) else (None, item)
            runs = ([(lead + "　", lead_color, True)] if lead else []) + [(body, SLATE, False)]
            # 英數夾在中文裡時，LibreOffice 常在整段英數前提早換行，可用寬度打九折估算
            h = need_height([nz((lead + "　" if lead else "") + body)], size, w * 0.90, sp) + 0.04
            self.tx(s, x, y, w, h, runs, size=size, sp=sp)
            y += h + gap
        return y

    def heading(self, s, x, y, w, text, size=14):
        self.tx(s, x, y, w, 0.30, text, size=size, color=RED, bold=True)
        return y + 0.36

    def table(self, s, x, y, widths, header, rows, size=12, hsize=12.5, pad=0.07, sp=1.08, center=(), hl_col=None):
        """紅底白字表頭＋交錯淡灰列，不畫格線。cell 可為字串或 (字串, 色, 粗體)。回傳結束 y。"""
        xs = [x]
        for wv in widths[:-1]:
            xs.append(xs[-1] + wv)
        tw = sum(widths)
        hh = max(0.34, max(need_height([nz(h_)], hsize, w_ - 2 * pad - 0.04, 1.0) for h_, w_ in zip(header, widths)) + 0.10)
        self.rr(s, x, y, tw, hh, RED, rad=0.12, name="thead")
        for k, h_ in enumerate(header):
            self.tx(s, xs[k] + pad, y, widths[k] - 2 * pad, hh, h_, size=hsize, color=WHITE, bold=True,
                    anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER if k in center else PP_ALIGN.LEFT)
        y += hh + 0.03
        for r, row in enumerate(rows):
            cells = [c if isinstance(c, tuple) else (c, None, None) for c in row]
            rh = max(0.32, max(need_height([nz(c[0])], size, (widths[k] - 2 * pad - 0.04) * 0.95, sp) for k, c in enumerate(cells)) + 2 * pad)
            if r % 2 == 0:
                self.rect(s, x, y, tw, rh, LGREY, name="zebra")
            for k, (txt, col, b) in enumerate(cells):
                self.tx(s, xs[k] + pad, y, widths[k] - 2 * pad, rh, txt, size=size,
                        color=col if col is not None else (CHAR if k == 0 else SLATE),
                        bold=b if b is not None else (k == 0), anchor=MSO_ANCHOR.MIDDLE, sp=sp,
                        align=PP_ALIGN.CENTER if k in center else PP_ALIGN.LEFT)
            y += rh
        return y

    def panel(self, s, x, y, w, title, items, fill=LGREY, size=12.5, title_color=RED_D, h_min=0):
        """淺灰（或紅暈）色底整塊，裡面只放文字：標題＋數段。回傳結束 y。"""
        inner = w - 0.4
        hs = [need_height([nz(t_)], size, inner * 0.90, 1.12) + 0.02 for t_ in items]
        h = max(h_min, 0.20 + 0.34 + sum(hs) + 0.06 * (len(items) - 1) + 0.16)
        self.rr(s, x, y, w, h, fill, rad=0.06, name="panel")
        self.tx(s, x + 0.2, y + 0.16, inner, 0.32, title, size=13.5, color=title_color, bold=True)
        yy = y + 0.52
        for t_, hh in zip(items, hs):
            self.tx(s, x + 0.2, yy, inner, hh, t_, size=size, color=SLATE, sp=1.12)
            yy += hh + 0.06
        return y + h

    def figure(self, s, png, top=1.42, bottom=6.05, maxw=12.0):
        from PIL import Image
        iw, ih = Image.open(png).size
        dh = bottom - top; dw = dh * iw / ih
        if dw > maxw:
            dw = maxw; dh = dw * ih / iw
        x = (13.333 - dw) / 2
        add_layers(s, Path(png).with_suffix(""), x, top, dw)     # 分層圖：放映時講到哪層出現哪層
        return top + dh

    def shot(self, s, name, cap, x, y, w, h):
        from PIL import Image
        self.rr(s, x, y, w, h, LGREY, rad=0.03, name="shotframe")
        img = SHOT / name
        iw, ih = Image.open(img).size
        dw = w - 0.3; dh = dw * ih / iw
        if dh > h - 0.62:
            dh = h - 0.62; dw = dh * iw / ih
        s.shapes.add_picture(str(img), I(x + (w - dw) / 2), I(y + 0.15), width=I(dw))
        self.tx(s, x + 0.15, y + h - 0.40, w - 0.3, 0.28, cap, size=12, color=GREY, align=PP_ALIGN.CENTER)


# 各頁內容
from pages_tech import build_pages  # noqa: E402


def build():
    d = Deck()
    build_pages(d, globals())
    for sl in d.prs.slides:
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        if sl.slide_layout.name != "頁首":
            animate_slide(sl)                          # 放映時內容與圖依序出現
    OUTDIR.mkdir(parents=True, exist_ok=True)
    out = OUTDIR / f"{NAME}.pptx"
    d.prs.save(str(out))
    qa_pptx(out, d.prs)
    return out, len(d.prs.slides)


BAD_CHARS = re.compile(r"[①-⓿㉑-㊿─-╿▀-▟■□▲▼◆◇★☆※\U0001F300-\U0001FAFF☀-⛿✀-➿]")
OK_FILLS = ("zebra", "cover-mask", "inner-status")


def qa_pptx(path, prs):
    bad = []
    W, H = I(13.333), I(7.5)
    for i, sl in enumerate(prs.slides, 1):
        boxes = []
        for sh in sl.shapes:
            if sh.left is None:
                continue
            l, t, w, h = sh.left, sh.top, sh.width or 0, sh.height or 0
            if l < 0 or t < 0 or l + w > W + I(0.01) or t + h > H + I(0.01):
                bad.append(f"S{i} 超出頁面：{sh.name}")
            is_fill = sh.shape_type == 1 and sh.fill.type == 1
            if is_fill and h <= I(0.1) and w > I(1):
                bad.append(f"S{i} 細色條：{sh.name}")
            if is_fill and sh.line.fill.type not in (None, 5):
                bad.append(f"S{i} 色塊帶描邊：{sh.name}")
            if sl.slide_layout.name == "內容" and (is_fill or sh.shape_type == 13) and sh.name != "cover-mask":
                if l + w > I(11.15) and t + h > I(6.40):
                    bad.append(f"S{i} 壓到右下logo：{sh.name}")
            if sh.shape_type == 17 and sl.slide_layout.name == "內容":       # 文字框也不准壓 logo
                if l + w > I(11.20) and t + h > I(6.45) and sh.text_frame.text.strip():
                    bad.append(f"S{i} 文字壓到右下logo：{sh.text_frame.text[:16]}")
            if is_fill and sh.name not in OK_FILLS:
                boxes.append((sh.name, l, t, l + w, t + h))
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                A, B = boxes[a], boxes[b]
                if A[1] < B[3] and B[1] < A[3] and A[2] < B[4] and B[2] < A[4]:
                    inside = (A[1] <= B[1] and A[2] <= B[2] and A[3] >= B[3] and A[4] >= B[4]) or \
                             (B[1] <= A[1] and B[2] <= A[2] and B[3] >= A[3] and B[4] >= A[4])
                    if not inside or "Oval" not in (A[0] + B[0]):
                        bad.append(f"S{i} 色塊重疊：{A[0]} 與 {B[0]}")
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not re.match(r"ppt/(slides|notesSlides)/\w+\.xml$", n):
                continue
            xml = z.read(n).decode("utf-8")
            if "/slides/" in n:
                if "<p:ph " in xml:
                    bad.append(f"{n} 殘留佔位框")
                for sz in re.findall(r'<a:rPr[^>]*\ssz="(\d+)"', xml):
                    if int(sz) < 1200:
                        bad.append(f"{n} 字級小於12：{int(sz) / 100}")
                if "<a:tbl>" in xml:
                    bad.append(f"{n} 用了原生表格（會帶格線）")
            for x in re.findall(r"<a:t>([^<]*)</a:t>", xml):
                if BAD_CHARS.search(x):
                    bad.append(f"{n} 圈碼、框線、裝飾符號或emoji：{x[:30]}")
                if "**" in x or "```" in x or re.match(r"^\s*(#|- |\* )", x):
                    bad.append(f"{n} markdown符號：{x[:30]}")
                if "雲碩" in x:
                    bad.append(f"{n} 公司名寫錯：{x[:30]}")
                if re.search(rf"{_CJK} +[A-Za-z0-9]|[A-Za-z0-9%] +{_CJK}", x):
                    bad.append(f"{n} 中英數之間有空白：{x[:40]}")
                if re.search(r"(?<!Zeabur)(8月|八月)(那次|事件)", x) and "Zeabur" not in x:
                    bad.append(f"{n} 提到事件卻沒寫主詞Zeabur：{x[:40]}")
    anim = check_timing(path)                          # 封面與 Thank You 以外每頁都要有進場動畫
    for i, sl in enumerate(prs.slides, 1):
        if sl.slide_layout.name != "頁首" and not anim.get(i):
            bad.append(f"S{i} 沒有進場動畫")
    globals()["TIMING"] = anim
    bad += PROBLEMS
    if bad:
        msg = f"{path.name} 檢查未過：\n" + "\n".join(dict.fromkeys(bad))
        if KEEP_GOING:
            print(msg); return
        raise SystemExit(msg)
    print(f"檢查通過：{path.name}；動畫數（頁：數量）{anim}")


def to_pdf(pptx):
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(OUTDIR), str(pptx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return pptx.with_suffix(".pdf")


def contact_sheet(pdf):
    from PIL import Image, ImageDraw
    QA.mkdir(parents=True, exist_ok=True)
    tmp = QA / "_pages"
    tmp.mkdir(exist_ok=True)
    for f in tmp.glob("*.jpg"):
        f.unlink()
    subprocess.run(["pdftoppm", "-jpeg", "-r", "60", str(pdf), str(tmp / "p")], check=True)
    pages = sorted(tmp.glob("p-*.jpg"), key=lambda p: int(p.stem.split("-")[1]))
    ims = [Image.open(p) for p in pages]
    w, h = ims[0].size
    cols = 4; rows = math.ceil(len(ims) / cols); pad = 16
    sheet = Image.new("RGB", (cols * (w + pad) + pad, rows * (h + pad + 22) + pad), "#d9dce0")
    dr = ImageDraw.Draw(sheet)
    for k, im in enumerate(ims):
        x = pad + (k % cols) * (w + pad); y = pad + (k // cols) * (h + pad + 22)
        dr.text((x, y), f"{k + 1}", fill="#2d2d2d")
        sheet.paste(im, (x, y + 18))
    out = QA / f"{NAME}-總覽.jpg"
    sheet.save(out, quality=88)
    for f in pages:
        f.unlink()
    tmp.rmdir()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--skip-diagrams", action="store_true")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--keep-going", action="store_true", help="檢查未過也照樣轉PDF（只供排版迭代，不可交付）")
    a = ap.parse_args()
    global KEEP_GOING
    KEEP_GOING = a.keep_going
    if not a.skip_diagrams:
        subprocess.run([sys.executable, str(HERE / "diagrams_tech.py")], check=True)
    out, count = build()
    print(f"{out}（{count}頁）")
    if not a.no_pdf:
        pdf = to_pdf(out)
        print("PDF：", pdf)
        print("總覽：", contact_sheet(pdf))


if __name__ == "__main__":
    main()
