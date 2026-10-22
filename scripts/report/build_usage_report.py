#!/usr/bin/env python3
"""xConsole 服務使用報告產生器：把主控台 /api/admin/usage-report 的數字做成給主管看的 Word 月報。

用法：
  python3 scripts/report/build_usage_report.py --days 30 --password-file ~/.xconsole/report.pw
  python3 scripts/report/build_usage_report.py --from-json dist-report/usage-report.json   # 離線重建

流程：抓數字 → 畫圖（matplotlib）→ 組 docx（python-docx，微軟正黑體、品牌紅）
      → soffice 轉 PDF 取頁碼回填目錄 → 重建 → 驗證（民國紀年、Markdown 外洩、條列比例）。
所有數字由程式帶入、不手寫；估算值在文中寫明假設；沒設定的單價直接寫沒設定，不自己補。

連線一律驗證主控台的憑證：預設用 deploy_console.sh 帶回本機的 ~/.xconsole/console-ca/<IP>.pem，
換成正式憑證後改用系統信任的憑證庫。不提供「略過驗證」的選項。
密碼只從檔案讀（權限要 0600），不從指令列、不從環境變數。
"""
import argparse
import json
import os
import re
import ssl
import stat
import subprocess
import sys
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone
from http.cookiejar import CookieJar
from urllib.parse import urlparse

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from matplotlib import font_manager  # noqa: E402

BRAND = "#bf181f"; INK = "#16202a"; INK2 = "#3d4b56"; GREY = "#6b7a85"; MIST = "#f4f7f9"
CYAN = "#0891b2"; GREEN = "#2e7d32"; AMBER = "#d97706"; LIGHT = "#e6eaee"
TPE = timezone(timedelta(hours=8))
FONT = "微軟正黑體"


# ------------------------------------------------------------------ 資料

def fetch(base, user, pw_file, days, cacert=None):
    st = os.stat(pw_file)
    if st.st_mode & (stat.S_IRWXG | stat.S_IRWXO):
        raise SystemExit(f"{pw_file} 的權限太寬，其他人讀得到。請先 chmod 600")
    pw = open(pw_file, encoding="utf-8").read().strip()
    host = urlparse(base).hostname
    if not cacert:
        cand = os.path.expanduser(f"~/.xconsole/console-ca/{host}.pem")
        cacert = cand if os.path.exists(cand) else None
    # 兩種信任依據都會驗證，只是來源不同：先試部署時帶回的主控台憑證（暫用自簽），
    # 主控台換成正式憑證後那份就不對了，改用系統的憑證庫。不提供略過驗證的選項。
    tries = ([ssl.create_default_context(cafile=cacert)] if cacert else []) + [ssl.create_default_context()]
    last = None
    for ctx in tries:
        jar = CookieJar()
        op = urllib.request.build_opener(urllib.request.HTTPSHandler(context=ctx),
                                         urllib.request.HTTPCookieProcessor(jar))
        req = urllib.request.Request(base + "/api/login", method="POST",
                                     data=json.dumps({"username": user, "password": pw}).encode(),
                                     headers={"Content-Type": "application/json"})
        try:
            op.open(req, timeout=30).read()
            break
        except urllib.error.URLError as exc:
            if "CERTIFICATE_VERIFY_FAILED" not in str(exc):
                raise
            last = exc
    else:
        raise SystemExit(f"主控台的憑證驗證不過，已停止（不會略過驗證）：{last}")
    return json.load(op.open(f"{base}/api/admin/usage-report?days={days}", timeout=120))


def n(v, d=0):
    if v is None:
        return "—"
    return f"{v:,.{d}f}" if isinstance(v, float) else f"{v:,}"


def pct(v):
    return "—" if v is None else f"{v * 100:.1f}%"


# ------------------------------------------------------------------ 圖

def setup_font():
    for name in ["Microsoft JhengHei", "微軟正黑體", "PingFang TC", "Heiti TC", "Noto Sans CJK TC"]:
        if any(f.name == name for f in font_manager.fontManager.ttflist):
            plt.rcParams["font.family"] = name
            return name
    return None


def style_ax(ax):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#d9dee3")
    ax.tick_params(colors=GREY, labelsize=10, length=0)
    ax.xaxis.grid(True, color=LIGHT, linewidth=0.8)
    ax.set_axisbelow(True)


def make_charts(r, outdir):
    setup_font()
    paths = {}
    nodes = r["assets"]["nodes"]

    # 圖 1：各機池子用量（淡色＝總量、實色＝峰值、標平均）
    if nodes:
        fig, ax = plt.subplots(figsize=(7.2, 0.7 + 0.55 * max(2, len(nodes))))
        style_ax(ax)
        labels = [x["node"] for x in nodes]
        tot = [x["pool_total_gib"] or 0 for x in nodes]
        peak = [x["pool_used_peak_gib"] or 0 for x in nodes]
        ax.barh(labels, tot, color=LIGHT, height=0.55)
        bars = ax.barh(labels, peak, color=CYAN, height=0.55)
        ax.invert_yaxis()
        ax.set_xlim(0, max(tot + [1]) * 1.32)
        for b, x in zip(bars, nodes):
            ax.text((x["pool_total_gib"] or 0) + max(tot + [1]) * 0.02, b.get_y() + b.get_height() / 2,
                    f"峰值 {n(x['pool_used_peak_gib'], 1)}／平均 {n(x['pool_used_avg_gib'], 1)}／共 {n(x['pool_total_gib'], 0)} GiB",
                    va="center", fontsize=9.5, color=INK)
        ax.set_title("各機記憶體池用量（實色＝本期峰值）", loc="left", fontsize=12, color=INK, pad=10)
        fig.tight_layout()
        paths["pool"] = os.path.join(outdir, "fig1-pool.png")
        fig.savefig(paths["pool"], dpi=200)
        plt.close(fig)

    # 圖 2：地端各欄位處理的 token（一個序列、直接標數字）
    ps = r["build_vs_buy"]["per_slot"][:8]
    if ps:
        fig, ax = plt.subplots(figsize=(7.2, 0.7 + 0.5 * max(2, len(ps))))
        style_ax(ax)
        labels = [f"{x['node']} {x['slot']}" for x in ps]
        vals = [x["prompt"] for x in ps]
        bars = ax.barh(labels, vals, color=BRAND, height=0.55)
        ax.invert_yaxis()
        ax.set_xlim(0, max(vals + [1]) * 1.35)
        for b, x in zip(bars, ps):
            ax.text(b.get_width() + max(vals + [1]) * 0.015, b.get_y() + b.get_height() / 2,
                    f"{n(x['prompt'])} token／{n(x['requests'])} 次", va="center", fontsize=9.5, color=INK)
        ax.set_title("地端模型本期匯入的提示量", loc="left", fontsize=12, color=INK, pad=10)
        fig.tight_layout()
        paths["tokens"] = os.path.join(outdir, "fig2-tokens.png")
        fig.savefig(paths["tokens"], dpi=200)
        plt.close(fig)

    # 圖 3：系統記憶體最低值與底線（只畫切兩塊與多卡的機器）
    rn = [x for x in nodes if x["memory_model"] != "unified_single" and x["ram_free_min_gib"] is not None]
    if rn:
        fig, ax = plt.subplots(figsize=(7.2, 0.7 + 0.55 * max(2, len(rn))))
        style_ax(ax)
        labels = [x["node"] for x in rn]
        vals = [x["ram_free_min_gib"] for x in rn]
        cols = [AMBER if (x["ram_floor_gib"] and v < x["ram_floor_gib"]) else GREEN for x, v in zip(rn, vals)]
        bars = ax.barh(labels, vals, color=cols, height=0.55)
        ax.invert_yaxis()
        top = max(vals + [x["ram_floor_gib"] or 0 for x in rn] + [1]) * 1.4
        ax.set_xlim(0, top)
        for b, x in zip(bars, rn):
            ax.text(b.get_width() + top * 0.015, b.get_y() + b.get_height() / 2,
                    f"最低剩 {n(x['ram_free_min_gib'], 1)} GiB（底線 {n(x['ram_floor_gib'], 0)}）",
                    va="center", fontsize=9.5, color=INK)
        ax.set_title("系統記憶體本期最低剩餘量（琥珀色＝曾低於底線）", loc="left", fontsize=12, color=INK, pad=10)
        fig.tight_layout()
        paths["ram"] = os.path.join(outdir, "fig3-ram.png")
        fig.savefig(paths["ram"], dpi=200)
        plt.close(fig)

    # 圖 4：登入與准入結果（狀態色）
    p, a = r["people"], r["integrations"]["admission"]
    names = ["登入成功", "登入失敗", "載入嘗試", "被准入擋下", "強制略過"]
    vals = [p["logins_ok"], p["logins_failed"], a["attempts"], a["blocked"], a["forced"]]
    cols = [GREEN, AMBER, CYAN, AMBER, BRAND]
    fig, ax = plt.subplots(figsize=(7.2, 3.0))
    style_ax(ax)
    ax.xaxis.grid(False)
    ax.yaxis.grid(True, color=LIGHT, linewidth=0.8)
    bars = ax.bar(names, vals, color=cols, width=0.5)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height(), f"{v:,} 次", ha="center", va="bottom", fontsize=10, color=INK)
    ax.set_ylim(0, max(vals + [1]) * 1.3)
    ax.set_title("本期登入與記憶體准入結果", loc="left", fontsize=12, color=INK, pad=10)
    fig.tight_layout()
    paths["guard"] = os.path.join(outdir, "fig4-guard.png")
    fig.savefig(paths["guard"], dpi=200)
    plt.close(fig)
    return paths


# ------------------------------------------------------------------ Word 工具

def set_font(run, size=11, bold=False, color=None):
    # 英數字也用微軟正黑體：Segoe UI 在 Mac 上沒有，會退回襯線字
    run.font.name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), "Microsoft JhengHei")
    rf.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def shade(cell, hex_color):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcpr.append(shd)


def fix_widths(t, widths_cm):
    """欄寬要三處一起設（tblW、gridCol、每格 tcW），Word 與 LibreOffice 才都會照做。"""
    t.autofit = False
    tblpr = t._element.tblPr
    for e in tblpr.findall(qn("w:tblW")) + tblpr.findall(qn("w:tblLayout")):
        tblpr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:type"), "dxa")
    w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblpr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = t._element.find(qn("w:tblGrid"))
    for gc, cm in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(cm * 567)))
    for row in t.rows:
        for i, cm in enumerate(widths_cm):
            row.cells[i].width = Cm(cm)


def no_borders(table):
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def para(doc, text, size=11, bold=False, color=None, align=None, space_after=6, keep_next=False, indent=False):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size, bold, color)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True
    return p


def body(doc, text):
    return para(doc, text, 11, False, INK, WD_ALIGN_PARAGRAPH.JUSTIFY, 8, indent=True)


def heading(doc, text):
    p = para(doc, text, 16, True, INK, space_after=10, keep_next=True)
    p.paragraph_format.page_break_before = True
    return p


def table(doc, header, rows, widths, caption=None, first_col_bold=False):
    if caption:
        para(doc, caption, 10, True, INK2, WD_ALIGN_PARAGRAPH.CENTER, 4, keep_next=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, BRAND)
        c.text = ""
        set_font(c.paragraphs[0].add_run(txt), 10, True, "#ffffff")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            set_font(cells[i].paragraphs[0].add_run(str(v)), 10, first_col_bold and i == 0, INK)
            if ri % 2 == 1:
                shade(cells[i], MIST)
    fix_widths(t, widths)
    # 短表整張不拆頁
    if len(rows) <= 10:
        for row in t.rows[:-1]:
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figure(doc, path, caption):
    if not path:
        return
    doc.add_picture(path, width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(doc, caption, 9.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, 10)


def stat_tiles(doc, items):
    """六個數字：兩列三格的色底方塊，沒有邊線、沒有色條。"""
    t = doc.add_table(rows=3, cols=5)
    no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fix_widths(t, [5.3, 0.35, 5.3, 0.35, 5.3])
    for i, (label, value, note) in enumerate(items):
        c = t.rows[(i // 3) * 2].cells[(i % 3) * 2]
        shade(c, MIST)
        c.text = ""
        p = c.paragraphs[0]
        set_font(p.add_run(label), 9, False, GREY)
        p.paragraph_format.space_after = Pt(2)
        p2 = c.add_paragraph()
        set_font(p2.add_run(value), 15, True, BRAND)
        p2.paragraph_format.space_after = Pt(2)
        p3 = c.add_paragraph()
        set_font(p3.add_run(note), 8.5, False, INK2)
        p3.paragraph_format.space_after = Pt(6)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("云碩科技股份有限公司　機密文件　第 "), 8, False, GREY)
    for code in ("PAGE", "NUMPAGES"):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), code)
        rr = OxmlElement("w:r")
        tt = OxmlElement("w:t")
        tt.text = "1"
        rr.append(tt)
        fld.append(rr)
        p._p.append(fld)
        if code == "PAGE":
            set_font(p.add_run(" 頁，共 "), 8, False, GREY)
    set_font(p.add_run(" 頁"), 8, False, GREY)


# ------------------------------------------------------------------ 組文件

TOC = ["一、本期摘要", "二、誰在用、用了多少", "三、地端模型處理了多少", "四、設備投資有沒有用到",
       "五、對外連線與操作安全", "六、系統健康與治理", "七、方法與限制"]


def build_docx(r, figs, out, org, preparer, toc_pages=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Microsoft JhengHei"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    footer(doc)
    P, S, B, A, I, H = (r["people"], r["spend"], r["build_vs_buy"], r["assets"], r["integrations"], r["health"])
    per = r["period"]
    gen = per["to"][:10]

    # ---- 封面：眉標 → 主標 → 報告期間 → 資訊表 → 公司與日期
    p = para(doc, "服務使用報告", 12, True, BRAND, WD_ALIGN_PARAGRAPH.CENTER, 8)
    p.paragraph_format.space_before = Pt(150)
    para(doc, f"{org}　xConsole 模型主控台", 26, True, INK, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para(doc, f"報告期間　{per['from'][:10]} 至 {gen}（{per['days']} 天）", 12, False, INK2,
         WD_ALIGN_PARAGRAPH.CENTER, 40)
    table(doc, ["項目", "內容"], [
        ["文件名稱", f"{org} xConsole 服務使用報告"], ["版本", "1.0"], ["發行日期", gen],
        ["編製單位", preparer], ["資料來源", "xConsole 主控台的操作紀錄與每 5 分鐘取樣（程式自動彙整）"],
        ["文件密級", "機密　Confidential"]], [4, 12.6], first_col_bold=True)
    p = para(doc, "云碩科技股份有限公司", 11, True, INK, WD_ALIGN_PARAGRAPH.CENTER, 2)
    p.paragraph_format.space_before = Pt(60)
    para(doc, gen, 9.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, 0)

    # ---- 目錄（靜態表，頁碼第二次建置回填）
    p = para(doc, "目錄", 16, True, INK, space_after=10)
    p.paragraph_format.page_break_before = True
    table(doc, ["章節", "頁碼"], [[t, (toc_pages or {}).get(t, "")] for t in TOC], [13.6, 3])

    # ---- 一、本期摘要
    heading(doc, TOC[0])
    loc = B["local"]
    hi = H["items"]
    stat_tiles(doc, [
        ("登入（成功／失敗）", f"{n(P['logins_ok'])}／{n(P['logins_failed'])}", f"共 {n(P['distinct_users'])} 個帳號登入過"),
        ("操作次數", n(P["operations"]), "啟停、換模型、情境、憑證等"),
        ("地端模型處理請求", n(loc["requests"]), f"匯入 {n(loc['prompt'])} 個 token"),
        ("提示快取命中", pct(B["cache_hit_ratio"]), "重用的提示不必重算"),
        ("GPU 晶片用電", f"{n(S['gpu_energy_kwh'], 2)} 度", f"約 NT$ {n(S['est_ntd'])}" if S["est_ntd"] else "未設定電價"),
        ("待處理事項", str(len(hi)), H["summary"]),
    ])
    if per.get("note"):
        body(doc, f"本期資料不完整：{per['note']}。以下各章的數字只涵蓋有取樣的期間，不應直接拿來與完整月份比較。")
    nodes = A["nodes"]
    work = (f"地端模型處理了 {n(loc['requests'])} 個請求，其中 {pct(B['cache_hit_ratio'])} 的提示內容直接重用快取，不必重算。"
            if loc["requests"] else "取樣期間地端模型沒有處理到請求。")
    body(doc, f"本期 xConsole 共納管 {len(nodes)} 台機器，期間共有 {n(P['logins_ok'])} 次成功登入、"
              f"{n(P['operations'])} 次操作。" + work
              + ("本期沒有需要處理的異常。" if not hi else f"系統本身列出 {len(hi)} 項待處理事項，詳見第六章。"))
    body(doc, S["assumption"] + " " + (B.get("assumption") or (B.get("cloud_equivalent") or {}).get("assumption") or ""))

    # ---- 二、誰在用
    heading(doc, TOC[1])
    body(doc, "xConsole 只開放具有管理權限的帳號登入，所有啟動、停止、更換模型、切換情境與更換憑證的操作都會留下紀錄，"
              "紀錄內容是誰、何時、從哪台電腦、做了什麼、結果如何；密碼、私鑰與權杖不會寫進紀錄。下表依操作類型統計本期次數。")
    label = {"slot.start": "啟動模型", "slot.stop": "停止模型", "slot.restart": "重啟模型", "model.swap": "更換模型",
             "profile.apply": "切換情境", "all.stop": "全部停止", "tls.inspect": "檢查憑證", "tls.apply": "套用憑證",
             "tls.rollback": "還原憑證", "tls.csr": "產生 CSR", "tls.hsts": "HSTS 開關", "tls.hostnames": "修改網址",
             "tls.bootstrap": "重產暫用自簽", "trust.import": "信任代理憑證", "account.password": "變更密碼"}
    rows = [[label.get(k, k), n(v)] for k, v in P["operations_by_type"].items()] or [["本期沒有操作紀錄", "—"]]
    table(doc, ["操作類型", "次數"], rows, [12.6, 4], caption="表 2-1　本期操作次數")
    rows = [[o["user"], n(o["operations"])] for o in P["operators"]] or [["本期沒有操作紀錄", "—"]]
    table(doc, ["帳號", "操作次數"], rows, [12.6, 4], caption="表 2-2　各帳號的操作次數")

    # ---- 三、地端模型處理了多少
    heading(doc, TOC[2])
    if loc["requests"]:
        body(doc, f"本期地端模型共處理 {n(loc['requests'])} 個請求，匯入提示 {n(loc['prompt'])} 個 token、"
                  f"生成回答 {n(loc['generated'])} 個 token。其中 {n(loc['reused'])} 個 token 直接重用前一輪的快取，"
                  f"占匯入量的 {pct(B['cache_hit_ratio'])}。寫程式的用戶端每一輪送出的都是「同一份脈絡再加長一點」，"
                  "快取命中率越高，使用者等待第一個字的時間就越短。")
    else:
        body(doc, "取樣期間地端模型沒有處理到請求，因此本章沒有處理量可以報告。"
                  "處理量由各機代理持續累計引擎的請求記錄，主控台每 5 分鐘取一次，兩次之間的差就是那段時間的量；"
                  "取樣累積一段時間之後，本章就會有數字。")
    ce = B.get("cloud_equivalent")
    if ce:
        body(doc, f"若這些處理量改由雲端模型提供，以 {ce['model']} 估算約需 US$ {n(ce['usd'], 2)}"
                  + (f"（約 NT$ {n(ce['ntd'])}）" if ce.get("ntd") else "") + f"。{ce['assumption']}。"
                  "這是估算值，用途是讓設備投資有一個可比較的參考，不是實際節省的金額。")
    else:
        body(doc, "本期尚未在主控台設定雲端參考單價，因此只報地端處理量，不估算若改用雲端需要多少費用。"
                  "要納入比較，可在「設定 → 使用報告」填入參考模型的官方定價。")
    # 沒資料就不畫空表：一列「本期沒有紀錄」擠在窄欄裡斷行，看起來像壞掉
    if B["per_slot"]:
        rows = [[x["node"], x["slot"], n(x["requests"]), n(x["prompt"]), n(x["reused"]), n(x["generated"])]
                for x in B["per_slot"]]
        table(doc, ["機器", "欄位", "請求數", "匯入 token", "重用快取", "生成 token"], rows,
              [2.6, 2.4, 2.2, 3.2, 3.0, 3.2], caption="表 3-1　各欄位本期處理量")
    figure(doc, figs.get("tokens"), "圖 3-1　地端模型本期匯入的提示量")

    # ---- 四、設備投資有沒有用到
    heading(doc, TOC[3])
    below = [x for x in nodes if x["below_floor_samples"]]
    body(doc, f"本期每 5 分鐘取樣一次各機狀態，共 {sum(x['samples'] for x in nodes):,} 筆。"
              "圖 4-1 是每台機器記憶體池的峰值與平均；對統一記憶體機器而言，記憶體占用比運算單元的忙碌程度更能反映設備有沒有被用到，"
              "因為模型一旦載入就會長期占著記憶體。"
              + (f"本期有 {len(below)} 台機器的系統記憶體曾經低於安全底線，這些機器若再多載一個模型，有機會整台卡死，詳見圖 4-2 與第六章。"
                 if below else "本期沒有任何機器的系統記憶體低於安全底線。"))
    mm = {"unified_split": "統一記憶體（切兩塊）", "unified_single": "統一記憶體（單一池）", "discrete": "獨立顯存"}
    rows = [[x["node"], mm.get(x["memory_model"], "—"), pct(x["online_ratio"]),
             f"{n(x['pool_used_avg_gib'], 1)}／{n(x['pool_used_peak_gib'], 1)}",
             n(x["pool_total_gib"], 0), "—" if x["memory_model"] == "unified_single" else n(x["ram_free_min_gib"], 1),
             f"{n(x['slots_busy_avg'], 1)}／{x['slots_total']}"] for x in nodes] or [["本期沒有取樣", "", "", "", "", "", ""]]
    table(doc, ["機器", "記憶體架構", "在線", "池子平均／峰值", "池子總量", "系統記憶體最低", "運作中欄位"], rows,
          [2.2, 3.4, 1.8, 2.8, 1.9, 2.5, 2.0], caption="表 4-1　各機設備使用概況（單位 GiB）")
    figure(doc, figs.get("pool"), "圖 4-1　各機記憶體池用量")
    figure(doc, figs.get("ram"), "圖 4-2　系統記憶體本期最低剩餘量")

    # ---- 五、對外連線與操作安全
    heading(doc, TOC[4])
    t = I.get("tls") or {}
    body(doc, f"主控台只提供 HTTPS。目前使用的憑證由 {t.get('issuer', '—')} 簽發，"
              f"{t.get('validToTaipei', '—')}（台北時間）到期，剩 {t.get('daysLeft', '—')} 天"
              + ("；這是系統自己產生的暫用自簽憑證，使用者第一次連線會看到安全警告，正式上線前應換成憑證廠商或貴單位 CA 簽發的憑證。"
                 if t.get("selfSigned") else "。")
              + ("HSTS 已開啟。" if t.get("hsts") else "HSTS 尚未開啟。"))
    trusted = sum(1 for x in I["agents"] if x["trusted"])
    body(doc, f"主控台連到各台機器的代理時同樣走 HTTPS，並以部署時經 SSH 取回的該機憑證驗證對方，"
              f"憑證不符就拒絕連線。本期納管的 {len(I['agents'])} 台中，{trusted} 台已建立信任。"
              f"每一次載入模型之前都會先做記憶體准入試算：本期共 {n(I['admission']['attempts'])} 次載入嘗試，"
              f"{n(I['admission']['blocked'])} 次被擋下，{n(I['admission']['forced'])} 次由操作者強制略過。")
    figure(doc, figs.get("guard"), "圖 5-1　本期登入與記憶體准入結果")
    rows = [[x["node"], "HTTPS", "已信任" if x["trusted"] else "尚未信任"] for x in I["agents"]]
    table(doc, ["機器", "連線方式", "代理憑證"], rows, [6, 4, 6.6], caption="表 5-1　各機代理的連線狀態")

    # ---- 六、系統健康與治理
    heading(doc, TOC[5])
    if hi:
        body(doc, f"xConsole 每 5 分鐘取樣一次各機狀態，並在產生報告時逐項檢查。本期共有 {len(hi)} 項需要注意，列在下表。")
        table(doc, ["等級", "事項", "處理方式"], [[x["level"], x["what"], x["fix"]] for x in hi],
              [1.8, 7.6, 7.2], caption="表 6-1　待處理事項")
    else:
        body(doc, "xConsole 每 5 分鐘取樣一次各機狀態，並在產生報告時逐項檢查，本期無異常。")
    body(doc, "治理上建議兩件事。第一，人員異動時把帳號停用而不是刪除，歷史操作紀錄才歸得了戶。"
              "第二，統一記憶體機器的安全底線不要為了多載一個模型而調低；這類機器一旦記憶體耗盡，不是單一服務當掉，而是整台卡死、需要到現場重開。")

    # ---- 七、方法與限制
    heading(doc, TOC[6])
    body(doc, "本報告的數字全部由 xConsole 主控台的操作紀錄與定時取樣直接計算後帶入，不經人工整理，在主控台上都能逐筆查到。"
              "取樣間隔為 5 分鐘，兩次取樣之間若超過 15 分鐘沒有資料，那段時間不計入用電與在線統計。")
    body(doc, "用電量只含 GPU 晶片自己回報的功耗，不含 CPU、記憶體、硬碟、電源轉換損失與空調，整機實際用電會明顯較高。"
              "地端模型的處理量取自引擎的請求記錄，引擎重啟或記錄輪替時以累計值歸零後重新計算。"
              "記憶體占用反映的是模型有沒有載入、占了多少，不等於運算單元的忙碌程度。")
    p = para(doc, f"本文件由云碩科技股份有限公司（xCloudinfo）編製　{gen[:4]} 年 {int(gen[5:7])} 月",
             9.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, 0)
    p.paragraph_format.space_before = Pt(30)
    doc.save(out)


# ------------------------------------------------------------------ 頁碼回填與驗證

def to_pdf(docx_path, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   capture_output=True, text=True, timeout=240)
    return os.path.splitext(docx_path)[0] + ".pdf"


def pages_for(pdf):
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    pages = txt.split("\f")
    found = {}
    for h in TOC:
        for i, pg in enumerate(pages):
            if i >= 2 and h in pg:
                found[h] = str(i + 1)
                break
    return found


def verify(pdf):
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    problems = []
    if re.search(r"民國|中華民國", txt):
        problems.append("出現民國紀年")
    if "第零章" in txt:
        problems.append("出現「第零章」")
    if "**" in txt or "`" in txt:
        problems.append("Markdown 標記外洩")
    if re.search(r"[─│█┌┐└┘├┤▲▼◆★※①②③]", txt):
        problems.append("出現框線或圈碼字元")
    if "雲碩" in txt:
        problems.append("公司名寫成「雲碩」")
    bullets = len(re.findall(r"^\s*•", txt, re.M))
    paras = len([ln for ln in txt.splitlines() if len(ln.strip()) > 30])
    if bullets > paras / 2:
        problems.append(f"條列過多（{bullets} 條列 vs {paras} 段落行）")
    return problems, txt.count("\f")


def main():
    ap = argparse.ArgumentParser(description="產出 xConsole 服務使用報告（Word＋PDF）")
    ap.add_argument("--base", default="https://192.168.11.27:9400")
    ap.add_argument("--user", default="admin")
    ap.add_argument("--password-file", default=os.path.expanduser("~/.xconsole/report.pw"))
    ap.add_argument("--cacert", help="主控台憑證；不給時用 ~/.xconsole/console-ca/<IP>.pem 或系統憑證庫")
    ap.add_argument("--days", type=int, default=30)
    ap.add_argument("--from-json")
    ap.add_argument("--out", default="dist-report")
    ap.add_argument("--org", default="云碩科技")
    ap.add_argument("--preparer", default="云碩科技股份有限公司（xCloudinfo）")
    a = ap.parse_args()

    r = (json.load(open(a.from_json, encoding="utf-8")) if a.from_json
         else fetch(a.base, a.user, a.password_file, a.days, a.cacert))
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, "usage-report.json"), "w", encoding="utf-8") as fh:
        json.dump(r, fh, ensure_ascii=False, indent=1)
    figs = make_charts(r, a.out)
    stem = f"xConsole-服務使用報告-{r['period']['to'][:10].replace('/', '')}"
    docx_path = os.path.join(a.out, stem + ".docx")
    build_docx(r, figs, docx_path, a.org, a.preparer)
    pdf = to_pdf(docx_path, a.out)
    toc = pages_for(pdf)
    build_docx(r, figs, docx_path, a.org, a.preparer, toc_pages=toc)
    pdf = to_pdf(docx_path, a.out)
    problems, pages = verify(pdf)
    print(f"報告：{docx_path}\nPDF：{pdf}（{pages} 頁）\n目錄頁碼：{toc}")
    print("驗證：" + ("通過" if not problems else "；".join(problems)))
    return 1 if problems or len(toc) != len(TOC) else 0


if __name__ == "__main__":
    sys.exit(main())
